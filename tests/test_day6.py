"""Day 6 workflow checks.

Windows Notepad, Word, and a real screen capture are not started here.
The file rules, the plan, and a stand-in companion are.
"""

from __future__ import annotations

import asyncio
import base64
import threading
from dataclasses import replace
from pathlib import Path

import pytest

from docx import Document

from machaa.agents.planner import Action
from machaa.agents.planning import parse_research_request
from machaa.agents.workflow import (
    _run_step,
    can_run_together,
    copy_matches,
    run_independent,
    skeleton_actions,
)
from machaa.companion.actions import CompanionState, perform
from machaa.companion.policy import resolve_docx, title_matches_window
from machaa.core.errors import CompanionRefused
from machaa.core.guardian import Guardian
from machaa.core.service import TaskService
from machaa.llm.fixture import FixtureGemini
from machaa.main import create_app
from machaa.memory.recall import describe_last, is_recall
from fastapi.testclient import TestClient
from tests.fakes import ExplodingBrowser, FakeBrowser

FLOW = (
    "research modern 2BHK interiors, prepare 3 creative Instagram posts for "
    "Vistaara Interiors, create a DOCX with captions and licensed pictures, "
    "save it in Documents/Machaa, and open it."
)
RECALL = "what did you create in my last task?"
PAYMENT = "buy an image using my credit card"


def _docs(settings) -> Path:
    return settings.data_dir / "Documents" / "Machaa"


def _slow(settings):
    return replace(settings, task_timeout_seconds=20)


class WorkflowCompanion:
    def __init__(self, documents: Path) -> None:
        self.documents = documents
        self.calls: list[str] = []

    def configured(self) -> bool:
        return True

    def set_permission(self, enabled: bool) -> dict:
        return {"ok": True, "enabled": enabled}

    def call(self, action: str, arguments: dict) -> dict:
        self.calls.append(action)
        if action == "desktop.save_docx":
            path = resolve_docx(self.documents, arguments["filename"])
            data = base64.b64decode(arguments["content"])
            path.parent.mkdir(parents=True, exist_ok=True)
            if path.is_file() and path.read_bytes() == data:
                return {
                    "ok": True,
                    "detail": f"Saved {path.name} in the Machaa Documents folder.",
                    "verified": True,
                    "duplicate": True,
                    "path": str(path),
                }
            path.write_bytes(data)
            return {
                "ok": True,
                "detail": f"Saved {path.name} in the Machaa Documents folder.",
                "verified": True,
                "path": str(path),
            }
        if action == "desktop.open_docx":
            path = resolve_docx(self.documents, arguments["filename"])
            if not path.is_file():
                raise CompanionRefused("That file is not in the Machaa Documents folder.")
            return {
                "ok": True,
                "detail": f"Opened {path.name} with its registered app.",
                "verified": True,
                "path": str(path),
            }
        if action == "desktop.open_explorer":
            return {
                "ok": True,
                "detail": "Opened File Explorer in the Machaa Documents folder.",
                "verified": True,
            }
        if action == "desktop.capture_screen":
            path = self.documents / "machaa-screen.bmp"
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_bytes(b"BM")
            return {"ok": True, "detail": "Saved machaa-screen.bmp in the Machaa Documents folder.", "verified": True}
        if action == "desktop.open_notepad":
            return {"ok": True, "detail": "Opened Notepad.", "pid": 4242}
        if action == "desktop.type_notepad":
            return {"ok": True, "detail": "Typed into the Notepad Machaa opened."}
        if action == "desktop.switch_window":
            return {"ok": True, "detail": "Switched to the window Machaa opened.", "verified": True}
        if action == "desktop.organize_file":
            return perform(CompanionState(self.documents), action, arguments) | {"ok": True}
        raise CompanionRefused(action)


def _placeholder(query, dest, guardian):
    from machaa.images.licensed import ImageResult, _write_placeholder_png

    path = dest / "reference-placeholder.png"
    _write_placeholder_png(path)
    return ImageResult(placeholder=True, path=path, title="placeholder")


async def _until(service: TaskService, task_id: str, statuses: set[str]) -> dict:
    current = {}
    for _ in range(200):
        current = service.get(task_id)
        if current["status"] in statuses:
            return current
        await asyncio.sleep(0.05)
    raise AssertionError(current)


def test_founder_sentence_is_a_workflow_and_old_sentences_stay():
    plan = parse_research_request(FLOW)
    assert plan is not None
    assert plan.kind == "workflow"
    assert plan.brand == "Vistaara Interiors"
    assert plan.query == "modern 2BHK interiors"
    assert plan.platform == "Instagram"
    social = parse_research_request(
        "Machaa, prepare an Instagram content document for Vistaara Interiors "
        "about affordable premium-looking 2BHK interiors. Research design ideas, "
        "find a suitably licensed reference image, and create 3 captions with hashtags."
    )
    assert social is not None and social.kind == "social"
    document = parse_research_request("Write a Word document about simple 2BHK interior design ideas")
    assert document is not None and document.kind == "document"
    search = parse_research_request("Machaa, open Google and search for modern bedroom interiors.")
    assert search is not None and search.kind == "search"
    notepad = parse_research_request(
        "Machaa, write three Instagram caption ideas for Vistaara Interiors in Notepad."
    )
    assert notepad is not None and notepad.kind == "desktop"


def test_independent_steps_can_overlap_and_desktop_steps_cannot():
    plan = parse_research_request(FLOW)
    actions = skeleton_actions(plan)
    waves = []
    for action in actions:
        if not waves or waves[-1][0].group != action.group:
            waves.append([action])
        else:
            waves[-1].append(action)
    assert can_run_together(waves[0]) is True
    assert any(item.tool == "browser.search_google" for item in waves[0])
    assert any(item.tool == "image.find_licensed" for item in waves[0])
    assert can_run_together(waves[1]) is False
    desktop = [
        Action("desktop.open_docx", {"filename": "vistaara-interiors.docx"}, "Open", 3),
        Action("desktop.open_explorer", {}, "Folder", 3),
    ]
    assert can_run_together(desktop) is False


async def test_independent_jobs_overlap():
    order: list[str] = []

    async def first():
        order.append("a-start")
        await asyncio.sleep(0.05)
        order.append("a-end")
        return "a"

    async def second():
        order.append("b-start")
        await asyncio.sleep(0.05)
        order.append("b-end")
        return "b"

    await run_independent([first, second])
    assert order.index("b-start") < order.index("a-end")


def test_window_identity_and_approved_folders(tmp_path: Path):
    assert title_matches_window("notepad", "Untitled - Notepad", 5, 5) is True
    assert title_matches_window("notepad", "notes.txt - Notepad++", 5, 5) is False
    assert title_matches_window("notepad", "Untitled - Notepad", 5, 9) is False
    assert title_matches_window("explorer", "Machaa", 4, 4) is True
    assert title_matches_window("explorer", "Machaa", 4, 8) is False
    assert title_matches_window("document", "vistaara-interiors.docx - Word", 3, 3) is True
    documents = tmp_path / "Machaa"
    documents.mkdir()
    state = CompanionState(documents)
    state.permission = True
    payload = Document()
    payload.add_paragraph("Hook: one light")
    docx = tmp_path / "made.docx"
    payload.save(str(docx))
    content = base64.b64encode(docx.read_bytes()).decode("ascii")
    saved = perform(state, "desktop.save_docx", {"filename": "vistaara-interiors.docx", "content": content})
    assert (documents / "vistaara-interiors.docx").is_file()
    again = perform(state, "desktop.save_docx", {"filename": "vistaara-interiors.docx", "content": content})
    assert again.get("duplicate") is True
    assert copy_matches(documents / "vistaara-interiors.docx", docx.read_bytes())
    moved = perform(
        state,
        "desktop.organize_file",
        {"filename": "vistaara-interiors.docx", "folder": "Posts"},
    )
    assert (documents / "Posts" / "vistaara-interiors.docx").is_file()
    assert "Posts" in moved["detail"]
    assert not (documents / "vistaara-interiors.docx").exists()
    with pytest.raises(CompanionRefused):
        perform(state, "desktop.open_docx", {"filename": "missing.docx"})
    with pytest.raises(CompanionRefused):
        perform(state, "desktop.save_docx", {"filename": "../secret.docx", "content": content})
    assert not (tmp_path / "secret.docx").exists()
    with pytest.raises(CompanionRefused):
        perform(state, "desktop.open_docx", {"filename": "Posts"})
    perform(state, "desktop.save_docx", {"filename": "vistaara-interiors.docx", "content": content})
    with pytest.raises(CompanionRefused) as opened:
        perform(state, "desktop.open_docx", {"filename": "vistaara-interiors.docx"})
    assert "Windows" in str(opened.value)
    assert "pretend" in str(opened.value)


def test_opening_a_document_on_this_machine_does_not_pretend(tmp_path: Path):
    documents = tmp_path / "Machaa"
    documents.mkdir()
    state = CompanionState(documents)
    state.permission = True
    payload = Document()
    payload.add_paragraph("Hello")
    made = tmp_path / "made.docx"
    payload.save(str(made))
    content = base64.b64encode(made.read_bytes()).decode("ascii")
    perform(state, "desktop.save_docx", {"filename": "vistaara-interiors.docx", "content": content})
    with pytest.raises(CompanionRefused) as caught:
        perform(state, "desktop.open_docx", {"filename": "vistaara-interiors.docx"})
    assert "Windows" in str(caught.value)
    assert "pretend" in str(caught.value)
    with pytest.raises(CompanionRefused) as screen:
        perform(state, "desktop.capture_screen", {})
    assert "pretend" in str(screen.value)
    blocked = Guardian().review_instruction(PAYMENT)
    assert blocked.allowed is False
    assert blocked.category == "payment"
    bad = Guardian().review_tool_call(
        "desktop.save_docx",
        {"filename": "vistaara-interiors.docx", "content": base64.b64encode(b"not a docx").decode("ascii")},
    )
    assert bad.allowed is False


def test_task_memory_uses_the_stored_task(settings):
    from machaa.core.tasks import TaskStore

    store = TaskStore(settings.database_path)
    created = store.create("research modern 2BHK interiors")
    store.add_step(
        created.id,
        tool="desktop.save_docx",
        label="Save the Word document",
        arguments={"filename": "vistaara-interiors.docx"},
        status="completed",
    )
    store.finish_step(store.get(created.id).steps[-1].id, "completed", "Saved vistaara-interiors.docx in the Machaa Documents folder.")
    store.finish(created.id, status="completed", summary="Created vistaara-interiors.docx.")
    text = describe_last(store, "not-stored")
    assert "vistaara-interiors.docx" in text
    assert "Created vistaara-interiors.docx." in text
    assert is_recall(RECALL) is True
    store.close()
    client = TestClient(create_app(settings, ExplodingBrowser))
    response = client.post("/api/tasks", json={"instruction": RECALL})
    assert response.status_code == 201
    body = response.json()
    assert body["status"] == "completed"
    assert "vistaara-interiors.docx" in body["summary"]
    denied = client.post("/api/tasks", json={"instruction": PAYMENT})
    assert denied.json()["status"] == "blocked"


async def test_full_workflow_saves_and_remembers(settings, monkeypatch):
    monkeypatch.setattr("machaa.images.licensed.find_licensed_image", _placeholder)
    settings = _slow(settings)
    companion = WorkflowCompanion(_docs(settings))
    service = TaskService(settings, FakeBrowser, FixtureGemini(), companion)
    service.desktop_enabled = True
    created = await service.submit(FLOW)
    pending = await _until(service, created["id"], {"needs_approval", "completed", "failed"})
    assert pending["status"] == "needs_approval", pending
    assert "Approve" in pending["summary"]
    service.approve(created["id"])
    finished = await _until(service, created["id"], {"completed", "failed", "blocked", "stopped"})
    assert finished["status"] == "completed", finished["summary"]
    assert "Opened" in finished["summary"]
    assert "placeholder" in finished["summary"].lower()
    assert "vistaara-interiors.docx" in finished["summary"]
    saved = _docs(settings) / "vistaara-interiors.docx"
    assert saved.is_file()
    text = "\n".join(paragraph.text for paragraph in Document(str(saved)).paragraphs)
    assert "Hook:" in text
    assert "Hashtags:" in text
    assert companion.calls.count("desktop.save_docx") == 1
    assert "desktop.open_docx" in companion.calls
    assert "os.system" not in companion.calls
    assert finished["saved_location"]
    recalled = await service.submit(RECALL)
    assert recalled["status"] == "completed"
    assert "vistaara-interiors.docx" in recalled["summary"]
    summary = await service.submit("Machaa, open Notepad and write a summary of that document.")
    typed = await _until(service, summary["id"], {"needs_approval", "completed", "failed"})
    assert typed["status"] == "needs_approval"
    service.approve(summary["id"])
    done = await _until(service, summary["id"], {"completed", "failed"})
    assert done["status"] == "completed", done["summary"]
    assert "summary" in done["summary"].lower()


async def test_stop_mid_task_does_not_claim_the_document_opened(settings, monkeypatch):
    monkeypatch.setattr("machaa.images.licensed.find_licensed_image", _placeholder)

    class GateBrowser(FakeBrowser):
        started = threading.Event()
        release = threading.Event()

        async def search_google(self, query: str) -> dict:
            GateBrowser.started.set()
            while not GateBrowser.release.is_set():
                await asyncio.sleep(0.05)
            return await super().search_google(query)

    GateBrowser.started.clear()
    GateBrowser.release.clear()
    settings = _slow(settings)
    companion = WorkflowCompanion(_docs(settings))
    service = TaskService(settings, GateBrowser, FixtureGemini(), companion)
    service.desktop_enabled = True
    created = await service.submit(FLOW)
    for _ in range(100):
        if GateBrowser.started.is_set():
            break
        await asyncio.sleep(0.05)
    assert GateBrowser.started.is_set()
    service.stop(created["id"])
    GateBrowser.release.set()
    finished = await _until(service, created["id"], {"stopped", "failed", "completed"})
    assert finished["status"] == "stopped", finished
    assert "Stopped" in finished["summary"]
    assert "did not mark the task done" in finished["summary"]
    assert "desktop.open_docx" not in companion.calls
    assert not (_docs(settings) / "vistaara-interiors.docx").exists()


async def test_a_rejected_save_writes_nothing(settings, monkeypatch):
    monkeypatch.setattr("machaa.images.licensed.find_licensed_image", _placeholder)
    settings = _slow(settings)
    companion = WorkflowCompanion(_docs(settings))
    service = TaskService(settings, FakeBrowser, FixtureGemini(), companion)
    service.desktop_enabled = True
    created = await service.submit(FLOW)
    pending = await _until(service, created["id"], {"needs_approval", "failed"})
    assert pending["status"] == "needs_approval"
    service.reject(created["id"])
    finished = await _until(service, created["id"], {"stopped", "failed", "completed"})
    assert finished["status"] == "stopped"
    assert "did not change anything" in finished["summary"].lower()
    assert "desktop.save_docx" not in companion.calls


async def test_unsafe_planner_output_is_not_run(settings, monkeypatch):
    monkeypatch.setattr("machaa.images.licensed.find_licensed_image", _placeholder)

    def evil(self, instruction, feedback="", max_steps=8):
        return [{"tool": "os.system", "arguments": {"cmd": "no"}, "label": "Run a shell", "group": 0}]

    monkeypatch.setattr(FixtureGemini, "plan_workflow", evil)
    settings = _slow(settings)
    companion = WorkflowCompanion(_docs(settings))
    service = TaskService(settings, FakeBrowser, FixtureGemini(), companion)
    service.desktop_enabled = True
    created = await service.submit(FLOW)
    pending = await _until(service, created["id"], {"needs_approval", "failed"})
    assert pending["status"] == "needs_approval", pending.get("summary")
    assert "os.system" not in companion.calls
    service.reject(created["id"])


class _Memory:
    def __init__(self, steps):
        self.steps = steps
        self.failed: list[tuple[int, str]] = []

    def get(self, task_id):
        return type("Task", (), {"steps": self.steps, "stop_requested": False, "status": "running"})()

    def finish_step(self, step_id, status, detail):
        self.failed.append((step_id, status))


async def test_a_finished_step_is_not_sent_again():
    step = type("Step", (), {"tool": "browser.search_google", "status": "completed", "id": 7})()
    agent = type("Agent", (), {"store": _Memory([step]), "calls": 0})()

    async def _call(*_args):
        agent.calls += 1
        return {"hits": [{"title": "Room"}], "url": "https://example.com"}

    agent._call = _call
    result = await _run_step(
        agent,
        "task",
        None,
        None,
        Action("browser.search_google", {"query": "rooms"}, "Search", 0),
        Path("."),
    )
    assert result["duplicate"] is True
    assert agent.calls == 0


async def test_a_failed_search_is_retried_once_then_the_task_stops(settings, monkeypatch):
    monkeypatch.setattr("machaa.images.licensed.find_licensed_image", _placeholder)

    attempts = {"n": 0}

    async def boom(*_args, **_kwargs):
        attempts["n"] += 1
        raise RuntimeError("the search page did not load")

    monkeypatch.setattr("machaa.agents.research.perform_search", boom)
    settings = _slow(settings)
    companion = WorkflowCompanion(_docs(settings))
    service = TaskService(settings, FakeBrowser, FixtureGemini(), companion)
    service.desktop_enabled = True
    created = await service.submit(FLOW)
    finished = await _until(service, created["id"], {"failed", "completed", "stopped", "blocked"})
    assert finished["status"] == "failed", finished["summary"]
    assert attempts["n"] == 2
    assert "desktop.open_docx" not in companion.calls
    assert "desktop.save_docx" not in companion.calls


async def test_screenshot_waits_for_consent(settings):
    companion = WorkflowCompanion(_docs(settings))
    service = TaskService(_slow(settings), ExplodingBrowser, FixtureGemini(), companion)
    service.desktop_enabled = True
    created = await service.submit("Machaa, take a screenshot of my screen.")
    pending = await _until(service, created["id"], {"needs_approval", "failed", "completed"})
    assert pending["status"] == "needs_approval", pending
    assert companion.calls == []
    service.approve(created["id"])
    finished = await _until(service, created["id"], {"completed", "failed"})
    assert finished["status"] == "completed", finished["summary"]
    assert companion.calls == ["desktop.capture_screen"]


def test_dashboard_still_has_the_microphone_and_the_step_list():
    script = Path("machaa/web/app.js").read_text(encoding="utf-8")
    page = Path("machaa/web/index.html").read_text(encoding="utf-8")
    assert 'data-testid="talk"' in page
    assert 'data-testid="recording-stop"' in page
    assert "saved-file" in script
    assert "step-state" in script
    assert "dataset.current" in script
    assert "data-testid=\"stop\"" in script or "controlButton(\"stop\"" in script
    assert Path("START-MACHAA.bat").is_file()
    assert Path("STOP-MACHAA.bat").is_file()
