"""Save a Word document from a draft. The file path is chosen by Machaa."""

from __future__ import annotations

import json
from pathlib import Path

from machaa.core.errors import MachaaError


class DocumentError(MachaaError):
    """The draft could not be saved as a Word document."""


def write_brief(path: Path, title: str, body: str) -> None:
    """Write title, sections, and the source links Machaa already collected."""

    from docx import Document

    try:
        payload = json.loads(body)
    except json.JSONDecodeError as exc:
        raise DocumentError("Machaa could not save the Word document.") from exc
    if not isinstance(payload, dict):
        raise DocumentError("Machaa could not save the Word document.")
    sections = payload.get("sections")
    if not isinstance(sections, list) or not sections:
        raise DocumentError("Machaa could not save the Word document.")
    path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    document.add_heading(title.strip() or "Machaa brief", level=0)
    document.add_paragraph(
        "Machaa wrote this from public search results on this computer. "
        "The wording came from Gemini. The links below are the pages Machaa opened."
    )
    for section in sections[:8]:
        if not isinstance(section, dict):
            continue
        heading = " ".join(str(section.get("heading") or "").split())
        paragraphs = section.get("paragraphs")
        if not heading or not isinstance(paragraphs, list):
            continue
        document.add_heading(heading[:160], level=1)
        for paragraph in paragraphs[:6]:
            text = " ".join(str(paragraph).split())
            if text:
                document.add_paragraph(text[:1200])
    sources = payload.get("sources") if isinstance(payload.get("sources"), list) else []
    kept = []
    for source in sources[:8]:
        if not isinstance(source, dict):
            continue
        url = str(source.get("url") or "").strip()
        label = " ".join(str(source.get("title") or "").split()) or url
        if url.startswith("http://") or url.startswith("https://"):
            kept.append((label[:200], url[:500]))
    if kept:
        document.add_heading("Sources", level=1)
        for label, url in kept:
            document.add_paragraph(f"{label} — {url}")
    document.save(str(path))
